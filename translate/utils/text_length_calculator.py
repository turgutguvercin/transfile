import csv
import pymupdf
from docx import Document
import openpyxl
import os


def calculate_length(filepath):
    _, extension = os.path.splitext(filepath)
    extension = extension.lower()
    if extension == ".pdf":
        return calculate_length_pdf(filepath)

    elif extension == ".docx":
        return calculate_length_docx(filepath)

    elif extension ==".xlsx":
        return calculate_length_xlsx(filepath)

    elif extension == ".csv":
        return calculate_length_csv(filepath)


def calculate_length_pdf(filepath):
    doc = pymupdf.open(filepath)
    total_text = ""
    for page in doc:
        total_text += page.get_text()
    doc.close()
    return len(total_text)

def calculate_length_docx(filepath):
    doc = Document(filepath)
    total_text = ""
    for para in doc.paragraphs:
        total_text += para.text + " "
    
    return len(total_text)

def calculate_length_xlsx(filepath):
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    total_text = ""
    for sheet in workbook.worksheets:
      for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                total_text += str(cell.value) + " "
    return len(total_text)

def calculate_length_csv(filepath):
    total_text = ""
    with open(filepath, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            for cell in row:
                total_text += str(cell) + " "
    return len(total_text)
